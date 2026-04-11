import os
import json
import re
import pdfplumber
import docx as python_docx
import urllib.request
import urllib.error

ANTHROPIC_API_KEY = os.environ.get('ANTHROPIC_API_KEY', '')

# ─── Text Extraction ─────────────────────────────────────────────────────────

def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF using pdfplumber."""
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        raise Exception(f"Erreur extraction PDF: {str(e)}")
    return text

def extract_text_from_docx(file_path: str) -> str:
    """Extract text from Word document."""
    try:
        doc = python_docx.Document(file_path)
        text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        # Also extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    text += "\n" + row_text
        return text
    except Exception as e:
        raise Exception(f"Erreur extraction DOCX: {str(e)}")

def extract_text(file_path: str, file_type: str) -> str:
    """Extract text based on file type."""
    if file_type == 'pdf':
        return extract_text_from_pdf(file_path)
    elif file_type in ['docx', 'doc']:
        return extract_text_from_docx(file_path)
    else:
        raise Exception(f"Type de fichier non supporté: {file_type}")

# ─── Smart Chunking ───────────────────────────────────────────────────────────

KEYWORDS_FR = [
    # Financier
    "montant", "prix", "coût", "budget", "paiement", "financement", "tva", "hors taxe",
    "xof", "fcfa", "euros", "acompte", "tranche", "garantie", "caution", "retenue",
    # Parties
    "maître d'ouvrage", "maitre d'ouvrage", "maître d'œuvre", "maitre d'oeuvre",
    "entrepreneur", "titulaire", "groupement", "sous-traitant", "prestataire",
    # Délais
    "délai", "delai", "durée", "duree", "calendrier", "planning", "notification",
    "démarrage", "demarrage", "réception", "reception", "achèvement",
    # Administratif
    "marché", "marche", "contrat", "lot", "référence", "numero", "approbation",
    "dncmp", "appel d'offres", "gré à gré", "passation", "tribunal", "litige",
    # Pénalités
    "pénalité", "penalite", "retard", "résiliation", "resiliation", "sanction",
    # Localisation
    "région", "region", "commune", "localité", "localite", "zone", "site",
    # CCAP keywords
    "article", "clause", "stipulation", "ordre de service"
]

def smart_chunk(text: str, max_chars: int = 8000) -> str:
    """
    Intelligent chunking: extract only relevant paragraphs containing keywords.
    For 2000-page documents this reduces tokens by ~95%.
    """
    paragraphs = [p.strip() for p in text.split('\n') if p.strip()]
    relevant = []
    seen = set()

    for i, para in enumerate(paragraphs):
        para_lower = para.lower()
        is_relevant = any(kw in para_lower for kw in KEYWORDS_FR)

        if is_relevant and para not in seen:
            # Add context: 1 line before and after
            context_start = max(0, i - 1)
            context_end = min(len(paragraphs), i + 2)
            for j in range(context_start, context_end):
                if paragraphs[j] not in seen:
                    relevant.append(paragraphs[j])
                    seen.add(paragraphs[j])

    result = "\n".join(relevant)

    # If still too long, truncate intelligently
    if len(result) > max_chars:
        # Keep first 3000 chars (usually has header info) + last 2000 + middle keywords
        first_part = result[:3000]
        last_part = result[-2000:]
        middle = result[3000:-2000]
        # From middle, take most keyword-dense sections
        middle_lines = middle.split('\n')
        scored = []
        for line in middle_lines:
            score = sum(1 for kw in KEYWORDS_FR if kw in line.lower())
            if score > 0:
                scored.append((score, line))
        scored.sort(reverse=True)
        top_middle = "\n".join([l for _, l in scored[:50]])
        result = first_part + "\n...\n" + top_middle + "\n...\n" + last_part

    return result

# ─── Claude API Call ──────────────────────────────────────────────────────────

def call_claude_haiku(prompt: str) -> str:
    """Call Claude Haiku API directly via urllib (no anthropic package needed)."""
    if not ANTHROPIC_API_KEY:
        raise Exception("Clé API Anthropic manquante. Configurez ANTHROPIC_API_KEY dans .env")

    url = "https://api.anthropic.com/v1/messages"
    headers = {
        "x-api-key": ANTHROPIC_API_KEY,
        "anthropic-version": "2023-06-01",
        "content-type": "application/json"
    }

    body = json.dumps({
        "model": "claude-haiku-4-5",
        "max_tokens": 4096,
        "messages": [{"role": "user", "content": prompt}]
    }).encode('utf-8')

    req = urllib.request.Request(url, data=body, headers=headers, method='POST')

    try:
        with urllib.request.urlopen(req, timeout=60) as response:
            result = json.loads(response.read().decode('utf-8'))
            return result['content'][0]['text']
    except urllib.error.HTTPError as e:
        error_body = e.read().decode('utf-8')
        raise Exception(f"Erreur API Claude: {e.code} - {error_body}")
    except Exception as e:
        raise Exception(f"Erreur connexion Claude: {str(e)}")

# ─── Resume Generation ────────────────────────────────────────────────────────

RESUME_PROMPT_TEMPLATE = """Tu es un expert juridique et technique en marchés publics et contrats BTP au Mali.
Analyse le texte de contrat ci-dessous et extrais EXACTEMENT les informations demandées.
Si une information n'est pas trouvée dans le texte, écris "Non mentionné".
Réponds UNIQUEMENT en JSON valide, sans texte avant ou après.

TEXTE DU CONTRAT:
{contrat_text}

Extrais ces informations et retourne ce JSON exact (complète toutes les valeurs):
{{
  "presentation_generale": {{
    "objet_contrat": "description complète de la nature et objet du projet",
    "parties_prenantes": {{
      "maitre_ouvrage": "nom et description du maître d'ouvrage",
      "maitre_oeuvre": "nom et description du maître d'œuvre (si différent)",
      "entrepreneur": "nom de l'entreprise/groupement titulaire",
      "sous_traitants": "sous-traitants mentionnés ou Non mentionné"
    }},
    "contexte_objectifs": "pourquoi ce contrat est signé et besoins couverts",
    "perimetre": {{
      "lots": "numéro(s) de lot(s) concerné(s)",
      "localisation": "lieu/région/communes du projet",
      "type_travaux": "type de travaux: bâtiment/route/AEP/assainissement/etc",
      "delai_execution": "durée d'exécution des travaux",
      "date_demarrage": "date de démarrage effectif des travaux ou Non mentionné",
      "ordre_de_service": "mention de l'ordre de service ou Non mentionné"
    }}
  }},
  "clauses_principales": {{
    "informations_administratives": {{
      "numero_marche": "numéro ou référence du marché",
      "numero_approbation": "numéro d'approbation DNCMP ou autorité compétente",
      "type_marche": "Public ou Privé",
      "mode_passation": "appel d'offres ouvert/restreint/gré à gré/consultation",
      "date_notification": "date de notification à l'entreprise",
      "source_financement": "État malien/ONG/bailleur international/KfW/BM/etc",
      "tribunal_competent": "juridiction compétente en cas de litige"
    }},
    "conditions_financieres": {{
      "montant_ht": "montant hors taxes en XOF/FCFA",
      "montant_ttc": "montant TTC en XOF/FCFA",
      "taux_tva": "taux TVA applicable",
      "conditions_paiement": "modalités: avancement/tranches/acomptes",
      "revision_prix": "prix ferme ou révisable",
      "caution_bonne_execution": "pourcentage de la caution de bonne exécution",
      "retenue_garantie": "pourcentage de la retenue de garantie"
    }},
    "obligations_parties": {{
      "obligations_maitre_ouvrage": "principales obligations du maître d'ouvrage",
      "obligations_maitre_oeuvre": "principales obligations du maître d'œuvre",
      "obligations_entrepreneur": "principales obligations de l'entreprise"
    }},
    "suivi_reception": {{
      "supervision": "mécanismes de supervision et contrôle",
      "reception_provisoire": "conditions de réception provisoire",
      "reception_definitive": "conditions de réception définitive"
    }},
    "clauses_particulieres": {{
      "penalites_retard": "taux et modalités des pénalités de retard",
      "conditions_resiliation": "conditions de résiliation du contrat",
      "reglement_litiges": "procédure de règlement des litiges"
    }}
  }}
}}"""

def generate_resume(file_path: str, file_type: str) -> dict:
    """
    Main function: extract text → chunk → call Claude → return structured resume.
    """
    # Step 1: Extract text
    full_text = extract_text(file_path, file_type)

    if not full_text.strip():
        raise Exception("Impossible d'extraire le texte du document")

    # Step 2: Smart chunk (handles 2000+ page documents)
    chunked_text = smart_chunk(full_text, max_chars=12000)

    # Step 3: Call Claude Haiku
    prompt = RESUME_PROMPT_TEMPLATE.format(contrat_text=chunked_text)
    response = call_claude_haiku(prompt)

    # Step 4: Parse JSON response
    # Clean potential markdown code blocks
    clean = re.sub(r'```(?:json)?\s*', '', response).strip()
    clean = re.sub(r'```\s*$', '', clean).strip()

    try:
        resume_data = json.loads(clean)
    except json.JSONDecodeError:
        # Try to extract JSON from response
        json_match = re.search(r'\{[\s\S]+\}', clean)
        if json_match:
            resume_data = json.loads(json_match.group())
        else:
            raise Exception("Erreur lors de l'analyse du contrat par l'IA")

    return resume_data
